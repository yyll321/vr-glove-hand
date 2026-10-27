"""Build the report from its reviewable Markdown and actual project evidence.

Run using the Codex bundled Python (python-docx and Pillow). No Unity rendering
is invented: the terrain figure is data, and the optical screenshots are labelled
historical. DOCX rendering and page QA must be performed separately.
"""
from pathlib import Path
import json
import math
import re
import struct
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Documentation'
FIG = OUT / 'figures'
STEM = '水下临场感系统与水动力扩展组会汇报'
FONT = Path('C:/Windows/Fonts/msyh.ttc')
BLUE = '#173e59'


def font(size):
    return ImageFont.truetype(str(FONT), size)


def architecture():
    im = Image.new('RGB', (1600, 720), 'white')
    d = ImageDraw.Draw(im)
    nodes = [
        ((30, 55, 475, 190), '观察与动作', ['XR 世界位置  右手柄  手套']),
        ((30, 285, 475, 420), '共同流场', ['背景扰动 → 体积网格 + 喷流']),
        ((30, 515, 475, 650), '接触与节流', ['手部近底扰动  尾流演示']),
        ((610, 55, 1030, 190), '世界与双眼几何', ['稳定观察  每眼深度重建']),
        ((610, 285, 1030, 420), '速度与局部浓度', ['输运  扩散  沉降  消光']),
        ((610, 515, 1030, 650), '声源与传播路径', ['分层折射反射  路径延迟']),
        ((1160, 55, 1570, 190), '水面与焦散', ['世界接收  共同波场']),
        ((1160, 285, 1570, 420), '颗粒 海草 灯光', ['局部流动  浑浊散射']),
        ((1160, 515, 1570, 650), '声音与接触反馈', ['环境声  ROV 声  触觉请求']),
    ]
    for box, title, lines in nodes:
        d.rounded_rectangle(box, 12, fill='#f2f7fa', outline='#90aabc', width=2)
        d.text((box[0]+20, box[1]+20), title, font=font(29), fill='black')
        for i, line in enumerate(lines):
            d.text((box[0]+20, box[1]+72+i*30), line, font=font(22), fill='#344f62')

    def arrow(a, b):
        d.line((a, b), fill=BLUE, width=4)
        angle=math.atan2(b[1]-a[1], b[0]-a[0])
        tip=[b, (b[0]-17*math.cos(angle-.4),b[1]-17*math.sin(angle-.4)),
             (b[0]-17*math.cos(angle+.4),b[1]-17*math.sin(angle+.4))]
        d.polygon(tip,fill=BLUE)
    for y in (123, 353, 583):
        arrow((475,y),(610,y));arrow((1030,y),(1160,y))
    arrow((820,190),(820,285))
    im.save(FIG/'system-architecture.png')


def terrain():
    text=(ROOT/'Assets/Underwater/Materials/SeabedTerrain_HydrodynamicRelief.asset').read_text(encoding='utf-8')
    data=bytes.fromhex(re.search(r'_typelessdata: ([0-9a-f]+)',text)[1])
    heights=[struct.unpack_from('<f',data,i*48+4)[0] for i in range(20301)]
    lo,hi=min(heights),max(heights)
    grid=Image.new('RGB',(101,201))
    for z in range(201):
        for x in range(101):
            t=(heights[z*101+x]-lo)/(hi-lo)
            grid.putpixel((x,200-z),(int(233-130*t),int(239-66*t),int(229-86*t)))
    im=Image.new('RGB',(1600,740),'white');d=ImageDraw.Draw(im)
    # Rotate the plan to keep the long Z axis horizontal on a portrait report.
    im.paste(grid.transpose(Image.Transpose.ROTATE_270).resize((960,480),Image.Resampling.BILINEAR),(60,95))
    def pos(x,z):return (60+(z+25)/50*960,95+(x+12.5)/25*480)
    points=[('ROVp',0,-2),('初始观察',0,-8),('管架',-5,8),('低通道',6,7),('标记',6,12),('巡检横管',0,15)]
    for n,(label,x,z) in enumerate(points,1):
        px,py=pos(x,z);d.ellipse((px-19,py-19,px+19,py+19),fill=BLUE)
        d.text((px-8,py-18),str(n),font=font(26),fill='white')
        d.text((1080,110+(n-1)*59),f'{n}  {label}',font=font(29),fill='black')
    d.rectangle((60,95,1020,575),outline=BLUE,width=2)
    d.text((60,30),'自然区局部地形 XZ 平面',font=font(32),fill='black')
    d.text((60,600),'横轴 Z 长 50 m    纵轴 X 宽 25 m',font=font(28),fill='black')
    d.text((60,650),f'高度 {lo:.2f} 至 {hi:.2f} m   20301 顶点   40000 三角形',font=font(27),fill='#344f62')
    d.text((1080,510),'ROV 平地半径 4 m',font=font(27),fill='black')
    d.text((1080,555),'出生平地半径 2.5 m',font=font(27),fill='black')
    d.text((1080,630),'保存顶点的数据示意',font=font(25),fill='black')
    im.save(FIG/'terrain-plan.png')


def visibility():
    im=Image.new('RGB',(1600,575),'white');d=ImageDraw.Draw(im)
    files=['11-seabed-low-visibility.png','12-seabed-upper-visibility.png']
    for i,(file,label) in enumerate(zip(files,['低处 底层浑浊较强','上层 视距逐渐恢复'])):
        image=Image.open(ROOT/'Logs/UnderwaterResearch'/file).convert('RGB')
        image.thumbnail((780,470))
        im.paste(image,(i*800+(800-image.width)//2,60+(470-image.height)//2))
        d.text((i*800+28,15),label,font=font(28),fill='black')
    d.text((28,535),'历史原研究场景桌面截图 不含本轮新增地形',font=font(24),fill='black')
    im.save(FIG/'historical-visibility.png')


def set_font(style, size, bold=False):
    style.font.name='Microsoft YaHei'
    style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')


def field(p, name):
    e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),name);p._p.append(e)


def table(doc, rows):
    tb=doc.add_table(rows=1, cols=len(rows[0]))
    tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    n=len(rows[0]); widths={2:[1.48,5.02],3:[1.25,2.8,2.45],4:[1.15,1.55,1.55,2.25]}[n]
    for col,w in zip(tb.columns,widths):col.width=Inches(w)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for idx,row in enumerate(rows):
        cells=tb.rows[0].cells if idx==0 else tb.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFEAF2' if idx==0 else ('F6F8FA' if idx%2==0 else 'FFFFFF'));tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.08
            if idx==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(value);run.font.size=Pt(10);run.font.bold=idx==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def build_doc():
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.64);s.bottom_margin=Inches(.58);s.left_margin=Inches(.82);s.right_margin=Inches(.82)
    s.header_distance=Inches(.27);s.footer_distance=Inches(.3)
    for name,size,bold in [('Normal',11,False),('Title',21,True),('Subtitle',12,False),('Heading 1',17,True),('Heading 2',12,True),('Caption',9.5,False)]:
        set_font(doc.styles[name],size,bold)
    for style in doc.styles:
        if style.type != 1:continue
        pr=style.element.get_or_add_pPr()
        for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)
    pstyle=doc.styles['Normal'].paragraph_format
    pstyle.line_spacing=1.12;pstyle.space_after=Pt(6)
    for name in ('Heading 1','Heading 2'):
        p=doc.styles[name].paragraph_format;p.space_before=Pt(10);p.space_after=Pt(7);p.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(14)
    doc.styles['Caption'].paragraph_format.space_after=Pt(8)
    h=s.header.paragraphs[0];h.add_run('水下临场感研究  技术实现与汇报材料').font.size=Pt(8)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('2026 10 26    ');field(foot,'PAGE');foot.add_run(' / ');field(foot,'NUMPAGES')
    for r in foot.runs:r.font.size=Pt(8)
    doc.core_properties.title=STEM;doc.core_properties.subject='Unity PICO 水下临场感系统技术报告'
    doc.core_properties.author='项目技术报告';doc.core_properties.keywords='水下; PICO; ROV; 流体; 临场感; 组会'
    lines=(OUT/(STEM+'.md')).read_text(encoding='utf-8').splitlines()
    i=0
    page_before=False
    references=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line=='<!-- pagebreak -->':page_before=True;continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            values=[[x.strip() for x in row.strip('|').split('|')] for row in rows if not re.match(r'^\|\s*---',row)]
            table(doc,values);continue
        match=re.match(r'!\[([^]]+)\]\(([^)]+)\)',line)
        if match:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(OUT/match[2]),width=Inches(6.45))
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title');continue
        if line.startswith('## '):
            references=line.startswith('## 19 ')
            p=doc.add_paragraph(line[3:],'Heading 1');p.paragraph_format.page_break_before=page_before and not references
            page_before=False;continue
        if line.startswith('### '):
            p=doc.add_paragraph(line[4:],'Heading 2')
            if line.startswith(('### 第 5 页','### 第 12 页')):p.paragraph_format.page_break_before=True
            continue
        style='Caption' if line.startswith(('图 1','图 2','图 3')) else 'Normal'
        p=doc.add_paragraph(line,style)
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.space_after=Pt(6)
            if references:p.paragraph_format.keep_together=True
    dest=OUT/(STEM+'.docx');doc.save(dest)
    print(dest)


def snapshot():
    historical=json.loads((ROOT/'Logs/UnderwaterResearch/results.json').read_text(encoding='utf-8'))
    checks=json.loads((ROOT/'Logs/PhysicsReview/scene-checks.json').read_text(encoding='utf-8'))
    saved=json.loads((ROOT/'Logs/PhysicsReview/offline-scene.json').read_text(encoding='utf-8'))
    data={
        'date':'2026-10-26','baseline_commit':'9fc1960',
        'historical_only':{k:historical[k] for k in ('passed','frameMs','causticCpuSubmitMs','gpu','limitation')},
        'historical_check_count':len(historical['checks']),
        'current_static_scene_checks':checks,
        'current_scene_sha256':saved['scene_sha256'],
        'current_terrain':saved['terrain'],
        'current_compile':'Runtime and Editor assemblies compiled with Unity Roslyn, five pre-existing warnings',
        'current_runtime':'NOT TESTED: Unity blocked by Sentinel LDK license dialog',
        'hmd_acceptance':'NOT TESTED for this extension; historical user confirmation of display/audio only',
    }
    (OUT/'validation-snapshot.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')


if __name__=='__main__':
    FIG.mkdir(parents=True,exist_ok=True)
    architecture();terrain();visibility();snapshot();build_doc()
